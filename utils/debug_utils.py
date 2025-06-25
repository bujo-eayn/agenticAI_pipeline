import json
import os
from pathlib import Path
from docx import Document
from docx.shared import Inches


def save_gpt_data_json(gpt_data: dict, filename: str = "debug_gpt_data01.json"):
    """Saves GPT data to a JSON file."""
    os.makedirs("debug", exist_ok=True)
    path = os.path.join("debug", filename)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(gpt_data, f, indent=4, ensure_ascii=False)
    return path


def create_structured_docx(gpt_data: dict, output_filename: str = "debug_structured_output01.docx"):
    """
    Creates a Word document based on GPT's structured output.
    Attempts to restore layout from chunked content.
    """
    os.makedirs("debug", exist_ok=True)
    doc = Document()

    doc.add_heading("🧠 AI-Reconstructed Document", 0)

    if "metadata" in gpt_data:
        meta = gpt_data["metadata"]
        doc.add_paragraph(f"📄 File Name: {meta.get('file_name', '')}")
        doc.add_paragraph(f"🧾 MIME Type: {meta.get('mime_type', '')}")
        doc.add_paragraph(f"📦 Size: {meta.get('file_size_bytes', 0)} bytes")
        doc.add_paragraph("")

    if "structured_output" in gpt_data:
        lines = gpt_data["structured_output"].splitlines()
        for line in lines:
            line = line.strip()
            if not line:
                continue
            if line.lower().startswith("heading") or line.lower().startswith("section"):
                doc.add_heading(line, level=2)
            elif line.lower().startswith("table:"):
                doc.add_paragraph(line, style="Intense Quote")
            elif line.lower().startswith("image") or "[image" in line.lower():
                doc.add_paragraph("🖼 [Image Placeholder]", style="Quote")
            else:
                doc.add_paragraph(line)

    path = os.path.join("debug", output_filename)
    doc.save(path)
    return path
