# utils/doc_utils.py
from docx import Document
from fpdf import FPDF
import os
from pathlib import Path


def convert_to_pdf(file_path):
    # Dummy placeholder: assume PDF is already uploaded
    if file_path.endswith(".pdf"):
        return file_path
    pdf_path = file_path + ".converted.pdf"
    with open(pdf_path, "wb") as f:
        f.write(b"%PDF-1.4\n...mock content...")
    return pdf_path


def create_final_doc(extracted_data, file_path):
    
    doc = Document()
    doc.add_heading("Final Extracted Document", 0)
    doc.add_paragraph(str(extracted_data))

    # Always use the fixed "outputs" directory
    base_path = Path("outputs")

    # Create the directory if it does not exist
    os.makedirs(base_path, exist_ok=True)

    # Determine a name based on the uploaded file name if provided
    if file_path and file_path[0]:
        filename_stem = Path(file_path[0]).stem
    else:
        filename_stem = "final_output"

    final_path = base_path / f"{filename_stem}.final.docx"

    # Save document
    doc.save(final_path)

    return str(final_path)



def create_text_doc(content: str, file_name="output.txt"):
    os.makedirs("output", exist_ok=True)
    path = os.path.join("output", file_name)
    with open(path, "w", encoding="utf-8") as f:
        f.write(content)
    return path

# This module provides utility functions for document processing.
# It includes converting documents to PDF and creating a final document from extracted data.